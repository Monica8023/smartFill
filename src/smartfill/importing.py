"""CSV/XLSX profile import, validation, normalization, and redaction."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from smartfill.secrets import SecretStore, redact_value


class ImportValidationError(ValueError):
    """Raised when an uploaded dataset cannot be imported safely."""


CANONICAL_FIELDS = frozenset(
    {
        "account.username",
        "account.password",
        "person.fullName",
        "person.gender",
        "person.idType",
        "person.idNumber",
        "person.phone",
        "person.email",
        "person.address",
    }
)
SECRET_FIELDS = frozenset({"account.password", "person.idNumber", "person.phone"})
PASSWORD_COLUMN_ALIASES = frozenset({"password", "passwd", "pwd", "密码", "登录密码"})
IDENTITY_COLUMN_ALIASES = frozenset(
    {"idnumber", "idcard", "identitynumber", "身份证", "身份证号", "证件号码"}
)
PHONE_COLUMN_ALIASES = frozenset({"phone", "mobile", "phonenumber", "手机号", "手机号码"})


class ProfileRecord(BaseModel):
    """Normalized record whose sensitive values have been replaced by references."""

    model_config = ConfigDict(frozen=True)

    id: str = Field(default_factory=lambda: str(uuid4()))
    values: dict[str, str]
    secret_refs: dict[str, str]


class ImportPreview(BaseModel):
    """Sanitized upload preview plus normalized execution records."""

    filename: str
    total_rows: int
    headers: list[str]
    rows: list[dict[str, str]]
    records: list[ProfileRecord]
    warnings: list[str] = Field(default_factory=list)


class ProfileImporter:
    """Import tabular data without retaining plaintext secrets in results."""

    def __init__(self, secret_store: SecretStore, max_bytes: int = 5 * 1024 * 1024) -> None:
        self._secret_store = secret_store
        self._max_bytes = max_bytes

    def preview(
        self,
        filename: str,
        content: bytes,
        mapping: dict[str, str],
        *,
        preview_rows: int = 20,
        allowed_fields: set[str] | frozenset[str] | None = None,
        required_fields: set[str] | frozenset[str] | None = None,
        sensitive_fields: set[str] | frozenset[str] | None = None,
    ) -> ImportPreview:
        extension = Path(filename).suffix.lower()
        if extension not in {".csv", ".xlsx"}:
            raise ImportValidationError("Only CSV or XLSX files are supported")
        if len(content) > self._max_bytes:
            raise ImportValidationError("Uploaded file is too large")
        if not content:
            raise ImportValidationError("Uploaded file is empty")

        rows = self._read_csv(content) if extension == ".csv" else self._read_xlsx(content)
        if not rows:
            raise ImportValidationError("The uploaded file has no data rows")
        headers = list(rows[0])
        self._validate_mapping(
            mapping,
            headers,
            allowed_fields=allowed_fields,
            required_fields=required_fields,
        )
        protected_fields = SECRET_FIELDS | frozenset(sensitive_fields or ())

        records: list[ProfileRecord] = []
        sanitized_rows: list[dict[str, str]] = []
        for row_number, row in enumerate(rows, start=2):
            values: dict[str, str] = {}
            secret_refs: dict[str, str] = {}
            sanitized = {
                source: self._sanitize_preview_column(source, value)
                for source, value in row.items()
            }
            record_id = str(uuid4())

            for source, canonical in mapping.items():
                raw_value = row.get(source, "").strip()
                if not raw_value:
                    continue
                normalized = self._normalize(canonical, raw_value, row_number)
                if canonical in protected_fields:
                    secret_refs[canonical] = self._secret_store.put(
                        f"{record_id}/{canonical}", normalized
                    )
                    sanitized[source] = (
                        "********" if canonical == "account.password" else redact_value(normalized)
                    )
                else:
                    values[canonical] = normalized

            records.append(ProfileRecord(id=record_id, values=values, secret_refs=secret_refs))
            if len(sanitized_rows) < preview_rows:
                sanitized_rows.append(sanitized)

        return ImportPreview(
            filename=Path(filename).name,
            total_rows=len(rows),
            headers=headers,
            rows=sanitized_rows,
            records=records,
        )

    @staticmethod
    def _sanitize_preview_column(source: str, value: str) -> str:
        if not value:
            return value
        normalized_source = (
            source.strip().lower().replace("_", "").replace("-", "").replace(" ", "")
        )
        if normalized_source in PASSWORD_COLUMN_ALIASES:
            return "********"
        if normalized_source in IDENTITY_COLUMN_ALIASES | PHONE_COLUMN_ALIASES:
            return redact_value(value.strip())
        return value

    @staticmethod
    def _validate_mapping(
        mapping: dict[str, str],
        headers: list[str],
        *,
        allowed_fields: set[str] | frozenset[str] | None = None,
        required_fields: set[str] | frozenset[str] | None = None,
    ) -> None:
        approved = CANONICAL_FIELDS if allowed_fields is None else frozenset(allowed_fields)
        for source, canonical in mapping.items():
            if source not in headers:
                raise ImportValidationError(f"Source column does not exist: {source}")
            if canonical not in approved:
                raise ImportValidationError(f"Unsupported canonical field: {canonical}")
        if len(set(mapping.values())) != len(mapping):
            raise ImportValidationError("A canonical field can only be mapped once")
        missing = frozenset(required_fields or ()) - set(mapping.values())
        if missing:
            raise ImportValidationError(f"Missing field mappings: {sorted(missing)}")

    @staticmethod
    def _read_csv(content: bytes) -> list[dict[str, str]]:
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ImportValidationError("CSV must use UTF-8 encoding") from error
        reader = csv.DictReader(io.StringIO(decoded))
        if not reader.fieldnames:
            raise ImportValidationError("CSV header row is required")
        return [
            {str(key): "" if value is None else str(value) for key, value in row.items()}
            for row in reader
        ]

    @staticmethod
    def _read_xlsx(content: bytes) -> list[dict[str, str]]:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (OSError, ValueError) as error:
            raise ImportValidationError("Invalid XLSX workbook") from error
        try:
            sheet = workbook.active
            if sheet is None:
                raise ImportValidationError("XLSX workbook has no active worksheet")
            iterator = sheet.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                raise ImportValidationError("XLSX header row is required")
            headers = [str(value).strip() if value is not None else "" for value in header_row]
            if any(not header for header in headers):
                raise ImportValidationError("XLSX headers cannot be empty")
            rows: list[dict[str, str]] = []
            for raw_row in iterator:
                values = ["" if value is None else str(value) for value in raw_row]
                if any(value.strip() for value in values):
                    rows.append(dict(zip(headers, values, strict=False)))
            return rows
        finally:
            workbook.close()

    @staticmethod
    def _normalize(canonical: str, value: str, row_number: int) -> str:
        if canonical == "person.gender":
            aliases = {
                "男": "male",
                "male": "male",
                "m": "male",
                "女": "female",
                "female": "female",
                "f": "female",
                "其他": "other",
                "other": "other",
            }
            try:
                return aliases[value.strip().lower()]
            except KeyError as error:
                raise ImportValidationError(f"Invalid gender at row {row_number}") from error
        if canonical == "person.email" and ("@" not in value or value.startswith("@")):
            raise ImportValidationError(f"Invalid email at row {row_number}")
        if canonical == "person.phone" and not value.replace("+", "", 1).isdigit():
            raise ImportValidationError(f"Invalid phone at row {row_number}")
        return value.strip()
